import pymupdf
import pytesseract

from PIL import Image
from docx import Document

from app.chunker import chunk_pages


pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)


def extract_text_from_pdf(file_path):
    document = pymupdf.open(file_path)

    text = ""

    for page in document:
        text += page.get_text()

    document.close()

    return text


def extract_pages_from_pdf(file_path):
    document = pymupdf.open(file_path)

    pages = []

    for page_number, page in enumerate(document):

        text = page.get_text()

        pages.append({
            "text": text,
            "page": page_number + 1
        })

    document.close()

    return pages


def extract_text_from_docx(file_path):
    document = Document(file_path)

    text = ""

    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            text += paragraph.text + "\n"

    return text


def extract_pages_from_docx(file_path):
    document = Document(file_path)

    pages = []

    text = ""

    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            text += paragraph.text + "\n"

    if text.strip():

        pages.append({
            "text": text,
            "page": 1
        })

    return pages


def extract_text_with_ocr(file_path):

    document = pymupdf.open(file_path)

    text = ""

    for page in document:

        pixmap = page.get_pixmap()

        image = Image.frombytes(
            "RGB",
            [pixmap.width, pixmap.height],
            pixmap.samples
        )

        text += pytesseract.image_to_string(
            image
        )

        text += "\n"

    document.close()

    return text


def extract_text_from_pdf_auto(file_path):

    text = extract_text_from_pdf(
        file_path
    )

    if len(text.strip()) > 50:

        return text

    print(
        "Very little text found. Using OCR..."
    )

    return extract_text_with_ocr(
        file_path
    )


def extract_pages_auto(file_path):

    if file_path.lower().endswith(
        ".docx"
    ):

        return extract_pages_from_docx(
            file_path
        )

    if file_path.lower().endswith(
        ".pdf"
    ):

        pages = extract_pages_from_pdf(
            file_path
        )

        total_text = "".join(
            page["text"]
            for page in pages
        )

        if len(total_text.strip()) > 50:

            return pages

        print(
            "Very little text found. Using OCR..."
        )

        ocr_text = extract_text_with_ocr(
            file_path
        )

        return [
            {
                "text": ocr_text,
                "page": 1
            }
        ]

    raise ValueError(
        "Unsupported file type."
    )


if __name__ == "__main__":

    pages = extract_pages_auto(
        "uploads/sample.pdf"
    )

    pdf_chunks = chunk_pages(
        pages
    )

    print("\nDOCUMENT CHUNKS:")

    for i, chunk in enumerate(
        pdf_chunks[:10]
    ):

        print(
            f"\nChunk {i + 1}"
        )

        print(
            f"Page: {chunk['page']}"
        )

        print(
            f"Text: {chunk['text'][:200]}"
        )